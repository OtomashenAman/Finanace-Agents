"""Finance data uploads: CSV/XLSX table -> column mapping -> validated canonical records.

Only the parsed records (and the file's SHA-256) are kept, never the file.
"""

import csv
import io
import re
import types
import typing
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook
from pydantic import ValidationError

from src.app.common.text_extraction.base import ExtractionError, check_zip_archive
from src.app.common.text_extraction.delimited import decode_text
from src.app.common.text_extraction.spreadsheet import cell_text
from src.app.data_sources.records import RECORD_LABELS, RECORD_MODELS, Record, RecordType

SUPPORTED_EXTENSIONS = (".csv", ".xlsx")
MAX_REPORTED_ERRORS = 50
PREVIEW_ROWS = 5

# Header spellings (normalized: lowercase, letters/digits only) recognized per field
_SYNONYMS: dict[str, tuple[str, ...]] = {
    "external_id": ("id", "externalid", "recordid", "systemid"),
    "name": (
        "name",
        "vendor",
        "vendorname",
        "supplier",
        "suppliername",
        "accountname",
        "displayname",
        "customer",
        "customername",
        "client",
    ),
    "email": ("email", "emailaddress", "vendoremail"),
    "tax_id": ("taxid", "vat", "vatnumber", "gstin", "ein", "tin"),
    "active": ("active", "isactive", "status"),
    "doc_number": (
        "docnumber",
        "invoicenumber",
        "invoiceno",
        "invoice",
        "billnumber",
        "billno",
        "reference",
    ),
    "vendor_name": ("vendor", "vendorname", "supplier", "suppliername", "payee", "merchant"),
    "total": (
        "total",
        "amount",
        "totalamount",
        "invoicetotal",
        "billtotal",
        "grandtotal",
        "pototal",
    ),
    "tax_total": ("tax", "taxtotal", "taxamount", "vatamount", "gst"),
    "bill_date": ("date", "billdate", "invoicedate", "txndate"),
    "due_date": ("duedate", "paymentdue"),
    "po_number": ("po", "ponumber", "pono", "purchaseorder", "purchaseordernumber"),
    "account": ("account", "glaccount", "gl", "expenseaccount", "glcode", "category"),
    "received_amount": (
        "received",
        "receivedamount",
        "receiptamount",
        "grnamount",
        "goodsreceived",
    ),
    "status": ("status", "postatus"),
    "po_date": ("date", "podate", "orderdate", "txndate"),
    "number": ("number", "accountnumber", "accountno", "acctnum", "code", "glcode"),
    "account_type": ("type", "accounttype", "category"),
    "current_balance": ("currentbalance", "closingbalance", "bankbalance", "balance"),
    # customers, invoices, payments
    "customer_name": ("customer", "customername", "client", "clientname", "billto"),
    "balance": ("balance", "openbalance", "balancedue", "amountdue", "outstanding", "openamount"),
    "invoice_date": ("date", "invoicedate", "issuedate", "txndate"),
    "paid_date": ("paiddate", "datepaid", "paymentdate"),
    "line_items": ("lineitems", "items", "products", "services", "performanceobligations"),
    "service_start": ("servicestart", "servicestartdate", "startdate", "contractstart"),
    "service_end": ("serviceend", "serviceenddate", "enddate", "contractend"),
    "reference_number": (
        "reference",
        "referenceno",
        "referencenumber",
        "paymentreference",
        "paymentref",
        "ref",
        "checkno",
        "wirereference",
    ),
    "payment_date": ("date", "paymentdate", "receiveddate", "valuedate", "txndate"),
    "invoice_number": ("invoice", "invoicenumber", "invoiceno", "appliedto"),
    "applied_amount": ("applied", "appliedamount", "amountapplied"),
    "unapplied_amount": ("unapplied", "unappliedamount"),
    "memo": ("memo", "note", "notes", "remittance", "remittanceadvice", "description"),
    # expenses (card, cash, check)
    "transaction_id": ("transactionid", "txnid", "cardtransactionid", "transaction"),
    "expense_date": ("date", "expensedate", "transactiondate", "postingdate", "txndate"),
    "payment_type": ("paymenttype", "paymentmethod", "method"),
    "employee": ("employee", "employeename", "cardholder", "spender"),
    "has_receipt": ("receipt", "hasreceipt", "receiptattached", "receiptprovided"),
    # month-end account balances (trial balance) and subledger totals
    "account_name": ("account", "accountname", "glaccount", "name"),
    "account_number": ("accountnumber", "accountno", "acctnum", "number", "code", "glcode"),
    "period": ("period", "month", "closeperiod", "yearmonth", "periodend", "date"),
    "ledger": ("ledger", "subledger"),
    # profit and loss by month
    "business_unit": ("businessunit", "unit", "class", "department", "segment", "bu"),
    "basis": ("basis", "scenario", "version", "actualbudget"),
    "revenue": ("revenue", "income", "sales", "totalincome", "netsales"),
    "cogs": ("cogs", "costofsales", "costofgoodssold"),
    "operating_expenses": ("operatingexpenses", "opex", "expenses", "totalexpenses"),
    "operating_profit": ("operatingprofit", "operatingincome", "ebit", "netoperatingincome"),
}

_DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d-%b-%Y", "%d %b %Y", "%Y/%m/%d")


@dataclass
class Table:
    headers: list[str]
    rows: list[list[str]]


@dataclass
class ImportResult:
    records: list[Record] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    error_count: int = 0


def _normalize(header: str) -> str:
    return re.sub(r"[^a-z0-9]", "", header.lower())


# Sheet names (normalized) recognized per record type, besides its label in RECORD_LABELS
_SHEET_NAMES: dict[RecordType, tuple[str, ...]] = {
    RecordType.VENDOR: ("vendors", "vendor", "suppliers", "vendormaster"),
    RecordType.BILL: ("bills", "bill", "supplierinvoices", "apinvoices", "payables"),
    RecordType.PURCHASE_ORDER: ("purchaseorders", "purchaseorder", "pos"),
    RecordType.ACCOUNT: ("chartofaccounts", "accounts", "coa"),
    RecordType.CUSTOMER: ("customers", "customer", "clients"),
    RecordType.INVOICE: ("invoices", "customerinvoices", "salesinvoices", "arinvoices"),
    RecordType.PAYMENT: ("payments", "customerpayments"),
    RecordType.EXPENSE: ("expenses", "cardtransactions", "cardexpenses"),
    RecordType.ACCOUNT_BALANCE: ("accountbalances", "trialbalance", "balances"),
    RecordType.SUBLEDGER_TOTAL: ("subledgertotals", "subledger", "agingtotals"),
    RecordType.PERIOD_RESULT: ("profitandloss", "pl", "pnl", "incomestatement"),
}


def sheet_record_type(sheet: str) -> RecordType | None:
    """The record type a workbook sheet holds, by its name ("Bills", "P&L", ...)."""
    name = _normalize(sheet)
    for record_type, names in _SHEET_NAMES.items():
        if name in names or name == _normalize(RECORD_LABELS[record_type]):
            return record_type
    return None


def _table(raw: list[list[str]], max_rows: int) -> Table:
    """First non-empty row is the header."""
    rows = [row for row in raw if any(row)]
    if len(rows) < 2:
        raise ExtractionError("The file needs a header row and at least one data row.")
    headers = [h or f"Column {i + 1}" for i, h in enumerate(rows[0])]
    body = [(row + [""] * len(headers))[: len(headers)] for row in rows[1:]]
    if len(body) > max_rows:
        raise ExtractionError(f"Too many rows ({len(body):,}); the limit is {max_rows:,}.")
    return Table(headers, body)


def read_workbook(data: bytes, filename: str, max_rows: int) -> list[tuple[str, Table | None]]:
    """Every sheet of an XLSX file, in order: (sheet name, table), or None when a sheet has no
    header and data row."""
    if not filename.lower().endswith(".xlsx"):
        raise ExtractionError("Upload an .xlsx file with one sheet per type.")
    check_zip_archive(data, "XLSX")
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many types for broken files
        raise ExtractionError("This XLSX file is damaged and can't be read.") from exc
    sheets: list[tuple[str, Table | None]] = []
    try:
        for sheet in workbook.worksheets:
            raw = [[cell_text(v) for v in row] for row in sheet.iter_rows(values_only=True)]
            try:
                table: Table | None = _table(raw, max_rows)
            except ExtractionError as exc:
                if str(exc).startswith("Too many rows"):
                    raise ExtractionError(f"Sheet '{sheet.title}': {exc}") from exc
                table = None
            sheets.append((sheet.title, table))
    finally:
        workbook.close()
    return sheets


def read_table(data: bytes, filename: str, max_rows: int) -> Table:
    """First sheet (XLSX) or the file (CSV): first non-empty row is the header."""
    name = filename.lower()
    if name.endswith(".csv"):
        text = decode_text(data)
        if text is None:
            raise ExtractionError("This CSV file isn't readable text.")
        try:
            dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(text[:4096], ",;\t|")
        except csv.Error:
            dialect = csv.excel
        raw = [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)]
    elif name.endswith(".xlsx"):
        check_zip_archive(data, "XLSX")
        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:  # openpyxl raises many types for broken files
            raise ExtractionError("This XLSX file is damaged and can't be read.") from exc
        try:
            sheet = workbook.worksheets[0]
            raw = [[cell_text(v) for v in row] for row in sheet.iter_rows(values_only=True)]
        finally:
            workbook.close()
    else:
        raise ExtractionError("Upload a .csv or .xlsx file.")
    return _table(raw, max_rows)


def fields_for(record_type: RecordType) -> dict[str, bool]:
    """Mappable fields of a record type -> whether they're required."""
    model = RECORD_MODELS[record_type]
    return {
        name: info.is_required() and name != "external_id"
        for name, info in model.model_fields.items()
    }


def suggest_mapping(record_type: RecordType, headers: list[str]) -> dict[str, str | None]:
    """field -> header, by known spellings. Each header is used at most once."""
    used: set[str] = set()
    mapping: dict[str, str | None] = {}
    normalized = {_normalize(h): h for h in headers}
    for name in fields_for(record_type):
        match = next(
            (
                normalized[s]
                for s in (_normalize(name), *_SYNONYMS.get(name, ()))
                if s in normalized and normalized[s] not in used
            ),
            None,
        )
        mapping[name] = match
        if match:
            used.add(match)
    return mapping


def _base_type(annotation: Any) -> Any:
    if isinstance(annotation, types.UnionType) or typing.get_origin(annotation) is typing.Union:
        return next(a for a in typing.get_args(annotation) if a is not type(None))
    return annotation


def _parse_decimal(value: str) -> Decimal:
    cleaned = re.sub(r"[^\d.\-()]", "", value)
    negative = cleaned.startswith("(") and cleaned.endswith(")")
    try:
        number = Decimal(cleaned.strip("()"))
    except InvalidOperation as exc:
        raise ValueError(f"'{value}' is not an amount") from exc
    return -number if negative else number


def _parse_date(value: str) -> date:
    candidate = value[:10] if re.match(r"\d{4}-\d{2}-\d{2}T", value) else value
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(candidate, fmt).date()  # noqa: DTZ007 - date only
        except ValueError:
            continue
    raise ValueError(f"'{value}' is not a date (use YYYY-MM-DD or DD/MM/YYYY)")


def _parse_bool(value: str) -> bool:
    lowered = value.lower()
    if lowered in ("1", "true", "yes", "y", "active"):
        return True
    if lowered in ("0", "false", "no", "n", "inactive"):
        return False
    raise ValueError(f"'{value}' is not yes/no")


def _convert(annotation: Any, value: str) -> Any:
    base = _base_type(annotation)
    if base is Decimal:
        return _parse_decimal(value)
    if base is date:
        return _parse_date(value)
    if base is bool:
        return _parse_bool(value)
    return value


def build_records(
    record_type: RecordType, table: Table, mapping: dict[str, str | None]
) -> ImportResult:
    model = RECORD_MODELS[record_type]
    columns = {h: i for i, h in enumerate(table.headers)}
    result = ImportResult()

    def error(line: int, message: str) -> None:
        result.error_count += 1
        if len(result.errors) < MAX_REPORTED_ERRORS:
            result.errors.append(f"Row {line}: {message}")

    seen: set[str] = set()
    for offset, row in enumerate(table.rows):
        line = offset + 2  # 1-based, after the header row
        values: dict[str, Any] = {"external_id": f"row-{line}"}
        problems: list[str] = []
        for name, header in mapping.items():
            if not header or header not in columns:
                continue
            raw = row[columns[header]].strip()
            if not raw:
                continue
            try:
                values[name] = _convert(model.model_fields[name].annotation, raw)
            except ValueError as exc:
                problems.append(f"{name}: {exc}")
        if problems:
            error(line, "; ".join(problems))
            continue
        try:
            record = model.model_validate(values)
        except ValidationError as exc:
            details = "; ".join(
                f"{e['loc'][0] if e['loc'] else 'row'}: {e['msg']}"
                for e in exc.errors(include_url=False)
            )
            error(line, details)
            continue
        if record.external_id in seen:
            error(line, f"duplicate id '{record.external_id}'")
            continue
        seen.add(record.external_id)
        result.records.append(record)
    return result
