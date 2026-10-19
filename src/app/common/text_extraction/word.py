import io

from docx import Document

from src.app.common.text_extraction.base import (
    ExtractedText,
    ExtractionError,
    check_zip_archive,
)

CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def extract_docx(data: bytes) -> ExtractedText:
    check_zip_archive(data, "DOCX")
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises many types for broken files
        raise ExtractionError("This Word file is damaged and can't be read.") from exc

    parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    if not parts:
        raise ExtractionError("This Word document has no text.")
    return ExtractedText("\n".join(parts), "docx", CONTENT_TYPE)
