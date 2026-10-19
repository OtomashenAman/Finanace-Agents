import io

import pytest
from docx import Document
from openpyxl import Workbook

from src.app.common.text_extraction.base import ExtractionError
from src.app.common.text_extraction.registry import extract_text


def make_pdf(page_texts: list[str]) -> bytes:
    """Minimal valid PDF; an empty string makes a page without any text (like a scan)."""
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>"]
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(len(page_texts)))
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(page_texts)} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, text in enumerate(page_texts):
        content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode() if text else b""
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {5 + 2 * i} 0 R >>".encode()
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream")
    out = b"%PDF-1.4\n"
    offsets = []
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    ).encode()
    return out


def make_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Invoice"
    sheet.append(["Supplier", "Invoice", "Total"])
    sheet.append(["Acme Cloud Services", "INV-2048", 12500.0])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_docx() -> bytes:
    document = Document()
    document.add_paragraph("Purchase order PO-778 for Acme Cloud Services")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Authorized"
    table.rows[0].cells[1].text = "12,500.00"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_pdf_text_is_extracted_per_page() -> None:
    result = extract_text(make_pdf(["Invoice INV-2048 total 12500.00", "Page two terms"]), "a.pdf")
    assert result.extractor == "pdf"
    assert result.page_count == 2
    assert "--- Page 1 ---\nInvoice INV-2048 total 12500.00" in result.text
    assert "Page two terms" in result.text


def test_scanned_pdf_without_text_is_rejected() -> None:
    with pytest.raises(ExtractionError, match="no text layer"):
        extract_text(make_pdf(["", ""]), "scan.pdf")


def test_pdf_with_some_empty_pages_warns() -> None:
    result = extract_text(make_pdf(["Invoice INV-2048 total 12500.00", ""]), "a.pdf")
    assert result.warnings == ["No text found on page(s) 2 (scanned pages?)"]


def test_xlsx_values_are_extracted_per_sheet() -> None:
    result = extract_text(make_xlsx(), "invoice.xlsx")
    assert result.extractor == "xlsx"
    assert result.text == (
        "## Sheet: Invoice\nSupplier | Invoice | Total\nAcme Cloud Services | INV-2048 | 12500"
    )


def test_docx_paragraphs_and_tables_are_extracted() -> None:
    result = extract_text(make_docx(), "po.docx")
    assert result.extractor == "docx"
    assert result.text == "Purchase order PO-778 for Acme Cloud Services\nAuthorized | 12,500.00"


def test_csv_in_legacy_encoding_and_semicolons() -> None:
    data = "Lieferant;Betrag\nMüller GmbH;12500,00\n".encode("cp1252")
    result = extract_text(data, "remit.csv")
    assert result.extractor == "csv"
    assert result.text == "Lieferant | Betrag\nMüller GmbH | 12500,00"


def test_detection_uses_content_not_extension() -> None:
    # A PDF named .csv is still read as a PDF
    assert extract_text(make_pdf(["Invoice INV-2048 total 12500.00"]), "x.csv").extractor == "pdf"


@pytest.mark.parametrize(
    ("data", "name", "message"),
    [
        (b"\x89PNG\r\n\x1a\n" + b"\x00" * 64, "scan.png", "no OCR"),
        (b"\xff\xd8\xff\xe0" + b"\x00" * 64, "photo.jpg", "no OCR"),
        (b"MZ\x90\x00 binary", "tool.exe", "Unsupported file type"),
        (b"plain words", "notes.md", "Unsupported file type"),
        (b"", "empty.csv", "empty"),
    ],
)
def test_unsupported_files_are_rejected(data: bytes, name: str, message: str) -> None:
    with pytest.raises(ExtractionError, match=message):
        extract_text(data, name)


def test_zip_bomb_is_refused(monkeypatch: pytest.MonkeyPatch) -> None:
    from src.app.common.text_extraction import base

    monkeypatch.setattr(base, "MAX_UNCOMPRESSED_BYTES", 100)
    with pytest.raises(ExtractionError, match="expands to more than"):
        extract_text(make_xlsx(), "big.xlsx")
